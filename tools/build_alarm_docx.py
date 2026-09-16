#!/usr/bin/env python3
import re
import shutil
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / "docs/01-告警中心管理功能说明.md"
OUTPUT = SOURCE.with_suffix(".docx")
MODULE_TITLE = SOURCE.stem.split("-", 1)[-1]
MODULE_NAME = MODULE_TITLE.replace("功能说明", "")
ASSET_DIR = ROOT / "output" / f"docx-{SOURCE.stem[:2]}-assets"

BLUE = "1F4E78"
CYAN = "2A9DCE"
LIGHT = "EAF3F8"
PALE = "F5F8FA"
GRID = "B8C9D6"
TEXT = "243746"
MUTED = "667B8B"
ORANGE = "E58A2B"


def font_path(bold=False):
    candidates = [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Medium.ttc" if bold else "/System/Library/Fonts/STHeiti Light.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
    ]
    return next((p for p in candidates if Path(p).exists()), None)


def img_font(size, bold=False):
    p = font_path(bold)
    return ImageFont.truetype(p, size) if p else ImageFont.load_default()


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=10.5, bold=False, color=TEXT, italic=False):
    run.font.name = "Arial Unicode MS"
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), "Arial Unicode MS")
    fonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    fonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def wrap_text(draw, text, font, width):
    lines, current = [], ""
    for ch in text:
        trial = current + ch
        if draw.textbbox((0, 0), trial, font=font)[2] <= width or not current:
            current = trial
        else:
            lines.append(current)
            current = ch
    if current:
        lines.append(current)
    return lines


def draw_box(draw, xy, label, fill=LIGHT, outline=CYAN, font=None):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=14, fill="#" + fill, outline="#" + outline, width=3)
    font = font or img_font(22, True)
    lines = wrap_text(draw, label, font, x2 - x1 - 24)
    heights = [draw.textbbox((0, 0), ln, font=font)[3] for ln in lines]
    y = (y1 + y2 - sum(heights) - 5 * (len(lines) - 1)) / 2
    for ln, h in zip(lines, heights):
        w = draw.textbbox((0, 0), ln, font=font)[2]
        draw.text(((x1 + x2 - w) / 2, y), ln, font=font, fill="#" + TEXT)
        y += h + 5


def arrow(draw, start, end, color=MUTED):
    draw.line([start, end], fill="#" + color, width=4)
    import math
    ang = math.atan2(end[1] - start[1], end[0] - start[0])
    for delta in (2.55, -2.55):
        p = (end[0] + 13 * math.cos(ang + delta), end[1] + 13 * math.sin(ang + delta))
        draw.line([end, p], fill="#" + color, width=4)


def render_framework_diagram(diagram_no, out_path, title, code=""):
    """Render concise framework diagrams: 4-6 nodes, one clear reading path."""
    specs = {
        1: [
            ("用户访问", "管理端｜物业移动端｜第三方系统"),
            ("应用功能", "查询统计｜告警处置｜工单联动"),
            ("业务服务", "事件标准化｜状态流转｜消息通知"),
            ("集成适配", "SDK｜GB/T 28181｜SNMP / Modbus｜API"),
            ("现场与数据", "设备系统｜业务数据库｜附件与审计"),
        ],
        2: [
            ("事件来源", "设备告警｜物业上报｜第三方事件"),
            ("统一接入", "校验来源并转换为标准告警模型"),
            ("规则处理", "分类、重复事件识别与状态初始化"),
            ("业务处置", "提醒、确认、派单与现场反馈"),
            ("闭环归档", "关闭告警并保存动态、附件和审计记录"),
        ],
        3: [
            ("访问终端", "管理浏览器｜物业移动端｜第三方应用"),
            ("平台入口", "统一 API 网关与身份认证"),
            ("平台服务", "告警｜工单｜媒体｜通知｜配置"),
            ("基础支撑", "数据库｜对象存储｜日志审计｜协议网关"),
        ],
        4: [
            ("发现", "现场设备或子系统产生事件"),
            ("接入", "适配器转换为标准告警"),
            ("确认", "告警中心提醒，值班人员核实"),
            ("处置", "按需转工单并完成现场处理"),
            ("关闭", "确认恢复并归档全过程记录"),
        ],
        5: [
            ("发现", "物业巡查发现馆内事件"),
            ("上报", "填写事件、位置并上传现场图片"),
            ("受理", "生成告警编号并进入待确认队列"),
            ("处置", "确认事件，按需创建工单"),
            ("归档", "核验处理结果并关闭"),
        ],
        6: [
            ("查看概况", "关注今日告警与待确认数量"),
            ("定位事件", "打开待办或查询历史告警"),
            ("核实详情", "查看基本信息、附件和动态"),
            ("选择方式", "直接处理或转为工单"),
            ("跟踪结果", "查看现场反馈与恢复情况"),
            ("确认关闭", "完成核验并归档"),
        ],
        7: [
            ("设置条件", "编号、等级、类型、状态和日期"),
            ("组合检索", "系统按全部条件取交集"),
            ("查看结果", "列表与详情保持同步"),
            ("继续操作", "查看详情、调整条件或重置"),
        ],
        8: [
            ("事件接收", "设备或物业事件进入告警服务"),
            ("记录提醒", "保存告警并通知值班人员"),
            ("确认研判", "值班人员核实并更新状态"),
            ("协同处置", "工单服务分派现场任务"),
            ("结果回传", "处理人员提交结果与附件"),
            ("核验关闭", "告警服务保存完整审计记录"),
        ],
    }
    titles = {
        1: "告警中心分层逻辑架构",
        2: "告警事件处理框架",
        3: "部署与接口边界",
        4: "硬件设备告警主流程",
        5: "物业移动端上报主流程",
        6: "管理人员日常处置流程",
        7: "告警组合查询流程",
        8: "告警处置协同时序",
    }
    if MODULE_NAME == "告警中心管理" and diagram_no in specs:
        title = titles[diagram_no]
        stages = specs[diagram_no]
    else:
        labels = []
        for label in re.findall(r"\[([^\]]+)\]", code):
            label = re.sub(r"<br\s*/?>", " ", label).strip()
            if label and label not in labels:
                labels.append(label)
        stages = [(label, "统一配置、记录与权限管理") for label in labels[:6]]
        if not stages:
            stages = [("业务入口", "管理业务请求"), ("平台服务", "统一处理与记录"), ("适配接入", "协议与设备适配"), ("业务结果", "返回状态与审计记录")]
    width, height = 1800, 690
    im = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(im)
    d.text((60, 38), title, font=img_font(31, True), fill="#" + BLUE)
    d.line((60, 90, width - 60, 90), fill="#" + CYAN, width=4)

    count = len(stages)
    gap = 28
    usable = width - 140
    box_w = (usable - gap * (count - 1)) / count
    y1, y2 = 180, 455
    positions = []
    for i, (label, detail) in enumerate(stages):
        x1 = 70 + i * (box_w + gap)
        x2 = x1 + box_w
        positions.append((x1, y1, x2, y2))
        d.rounded_rectangle((x1, y1, x2, y2), radius=18, fill="#" + PALE, outline="#" + CYAN, width=3)
        d.rounded_rectangle((x1, y1, x2, y1 + 68), radius=18, fill="#" + BLUE, outline="#" + BLUE, width=2)
        # Cover the lower rounded corners of the header band.
        d.rectangle((x1, y1 + 45, x2, y1 + 68), fill="#" + BLUE)
        lf = img_font(23, True)
        tw = d.textbbox((0, 0), label, font=lf)[2]
        d.text(((x1 + x2 - tw) / 2, y1 + 18), label, font=lf, fill="white")
        df = img_font(20)
        lines = wrap_text(d, detail, df, box_w - 38)
        line_h = 31
        yy = y1 + 100 + max(0, (120 - len(lines[:4]) * line_h) / 2)
        for line in lines[:4]:
            tw = d.textbbox((0, 0), line, font=df)[2]
            d.text(((x1 + x2 - tw) / 2, yy), line, font=df, fill="#" + TEXT)
            yy += line_h
        if i < count - 1:
            arrow(d, (x2 + 5, (y1 + y2) / 2), (positions[-1][2] + gap - 6, (y1 + y2) / 2), color=BLUE)

    note = "框架图仅表达主路径；字段、规则、权限与异常处理详见对应章节文字说明。"
    nf = img_font(19)
    d.rounded_rectangle((300, 535, width - 300, 610), radius=14, fill="#" + LIGHT, outline="#D2E3ED", width=2)
    tw = d.textbbox((0, 0), note, font=nf)[2]
    d.text(((width - tw) / 2, 558), note, font=nf, fill="#" + MUTED)
    im.save(out_path, quality=95)


def parse_mermaid_nodes(code):
    nodes = {}
    for line in code.splitlines():
        for m in re.finditer(r"\b([A-Za-z][A-Za-z0-9]*)\s*(?:\[([^\]]+)\]|\{([^}]+)\}|\(\(([^)]+)\)\))", line):
            nodes[m.group(1)] = next((g for g in m.groups()[1:] if g), m.group(1)).replace("[(", "").replace(")]", "")
    edges = []
    for line in code.splitlines():
        ids = re.findall(r"\b([A-Za-z][A-Za-z0-9]*)\s*(?=(?:\[|\{|\(\())", line)
        if "-->" in line:
            raw = re.split(r"-->(?:\|[^|]+\|)?", line)
            chain = []
            for part in raw:
                m = re.search(r"\b([A-Za-z][A-Za-z0-9]*)", part.strip())
                if m and m.group(1) in nodes:
                    chain.append(m.group(1))
            edges.extend(zip(chain, chain[1:]))
    return nodes, edges


def render_flowchart(code, out_path, title):
    nodes, edges = parse_mermaid_nodes(code)
    labels = list(nodes.values()) or ["流程图"]
    n = len(labels)
    cols = 4 if n > 12 else 3 if n > 6 else min(4, n)
    rows = (n + cols - 1) // cols
    width = 1800
    box_w, box_h = 360, 105
    xgap = (width - cols * box_w) // (cols + 1)
    height = 135 + rows * 155 + 45
    im = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(im)
    d.text((55, 35), title, font=img_font(30, True), fill="#" + BLUE)
    d.line((55, 85, width - 55, 85), fill="#" + CYAN, width=4)
    positions = {}
    ids = list(nodes.keys())
    for i, ident in enumerate(ids):
        r, c0 = divmod(i, cols)
        c = c0 if r % 2 == 0 else cols - 1 - c0
        x = xgap + c * (box_w + xgap)
        y = 120 + r * 155
        positions[ident] = (x, y, x + box_w, y + box_h)
    # Keep edge drawing behind boxes. When graph parsing is incomplete, use reading-order links.
    use_edges = [(a, b) for a, b in edges if a in positions and b in positions]
    if not use_edges:
        use_edges = list(zip(ids, ids[1:]))
    for a, b in use_edges:
        aa, bb = positions[a], positions[b]
        arrow(d, ((aa[0] + aa[2]) / 2, aa[3]), ((bb[0] + bb[2]) / 2, bb[1]))
    for ident in ids:
        draw_box(d, positions[ident], nodes[ident], fill=PALE if "{" not in code else LIGHT)
    im.save(out_path, quality=95)


def render_sequence(code, out_path, title):
    participants = []
    for line in code.splitlines():
        m = re.search(r"participant\s+(\w+)\s+as\s+(.+)", line.strip())
        if m:
            participants.append((m.group(1), m.group(2)))
    messages = []
    for line in code.splitlines():
        m = re.search(r"(\w+)-+>>?(\w+):\s*(.+)", line.strip())
        if m:
            messages.append(m.groups())
    width = 1800
    height = 180 + max(1, len(messages)) * 78 + 80
    im = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(im)
    d.text((55, 32), title, font=img_font(30, True), fill="#" + BLUE)
    d.line((55, 82, width - 55, 82), fill="#" + CYAN, width=4)
    margin = 100
    step = (width - 2 * margin) / max(1, len(participants) - 1)
    xs = {p[0]: margin + i * step for i, p in enumerate(participants)}
    for ident, label in participants:
        x = xs[ident]
        d.rounded_rectangle((x - 105, 105, x + 105, 160), radius=10, fill="#" + LIGHT, outline="#" + CYAN, width=2)
        f = img_font(17, True)
        lines = wrap_text(d, label, f, 190)
        yy = 115
        for ln in lines[:2]:
            tw = d.textbbox((0, 0), ln, font=f)[2]
            d.text((x - tw / 2, yy), ln, font=f, fill="#" + TEXT)
            yy += 21
        d.line((x, 160, x, height - 40), fill="#A9BAC5", width=2)
    for i, (src, dst, label) in enumerate(messages):
        if src not in xs or dst not in xs:
            continue
        y = 205 + i * 78
        arrow(d, (xs[src], y), (xs[dst], y), color=BLUE)
        f = img_font(17)
        lines = wrap_text(d, label, f, abs(xs[dst] - xs[src]) - 20 if abs(xs[dst] - xs[src]) > 120 else 260)
        text_line = lines[0] if lines else label
        tw = d.textbbox((0, 0), text_line, font=f)[2]
        d.rectangle(((xs[src] + xs[dst]) / 2 - tw / 2 - 6, y - 27, (xs[src] + xs[dst]) / 2 + tw / 2 + 6, y - 4), fill="white")
        d.text(((xs[src] + xs[dst]) / 2 - tw / 2, y - 27), text_line, font=f, fill="#" + TEXT)
    im.save(out_path, quality=95)


def configure_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    for key in ("ascii", "hAnsi", "eastAsia"):
        normal._element.rPr.rFonts.set(qn(f"w:{key}"), "Arial Unicode MS")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 8),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 11.5, "315E7D", 10, 5),
    ]:
        st = styles[name]
        st.font.name = "Arial Unicode MS"
        for key in ("ascii", "hAnsi", "eastAsia"):
            st._element.rPr.rFonts.set(qn(f"w:{key}"), "Arial Unicode MS")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        st = styles[name]
        st.font.name = "Arial Unicode MS"
        for key in ("ascii", "hAnsi", "eastAsia"):
            st._element.rPr.rFonts.set(qn(f"w:{key}"), "Arial Unicode MS")
        st.font.size = Pt(10.5)
        st.paragraph_format.left_indent = Inches(0.38)
        st.paragraph_format.first_line_indent = Inches(-0.19)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.25

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(hp.add_run(f"体育博物馆智慧化综合管理平台｜{MODULE_NAME}"), 8.5, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(fp.add_run("第 "), 8.5, color=MUTED)
    add_field(fp, "PAGE")
    set_run_font(fp.add_run(" 页"), 8.5, color=MUTED)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(100)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("体育博物馆智慧化综合管理平台"), 15, True, MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(10)
    set_run_font(p.add_run(MODULE_TITLE), 28, True, BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("功能架构｜业务流程｜操作说明｜验收建议"), 12, color=CYAN)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(92)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("文档版本 V1.0"), 11, True, TEXT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("编制日期：2026年8月31日"), 10.5, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("验收版本：完整功能与配置能力"), 10.5, color=MUTED)
    doc.add_page_break()


def add_toc(doc, headings):
    p = doc.add_paragraph("目录", style="Heading 1")
    p.paragraph_format.space_after = Pt(14)
    for level, text in headings:
        if level > 1:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        set_run_font(p.add_run(text), 10.5, True, BLUE)
    doc.add_page_break()


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    usable = 9840
    # Short first/number columns stay compact; narrative columns receive remaining width.
    first = 1050 if cols >= 3 else 2200
    widths = [first] + [(usable - first) // (cols - 1)] * (cols - 1) if cols > 1 else [usable]
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell = table.cell(ri, ci)
            cell.width = Inches(widths[ci] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            text = row[ci].strip() if ci < len(row) else ""
            set_run_font(p.add_run(text), 8.4 if cols >= 4 else 9, ri == 0, "FFFFFF" if ri == 0 else TEXT)
            if ri == 0:
                set_cell_shading(cell, BLUE)
            elif ri % 2 == 0:
                set_cell_shading(cell, PALE)
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_picture(doc, path, width=6.55, alt_text="功能界面或流程示意图"):
    # Keep portrait-oriented UI details inside the printable body instead of
    # scaling their width to the page and clipping the lower controls.
    with Image.open(path) as source_image:
        pixel_width, pixel_height = source_image.size
    rendered_height = width * pixel_height / max(pixel_width, 1)
    if rendered_height > 7.05:
        width = 7.05 * pixel_width / pixel_height
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    shape = r.add_picture(str(path), width=Inches(width))
    # Word accessibility metadata: keep the Markdown caption/diagram title on
    # the embedded image so screen readers do not encounter anonymous images.
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text)


def parse_md(lines):
    headings = []
    for line in lines:
        m = re.match(r"^(#{2,4})\s+(.+)$", line)
        if m and not (m.group(2) == MODULE_TITLE):
            headings.append((len(m.group(1)) - 1, m.group(2)))
    return headings


def clean_inline(text):
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    return text.replace("`", "").replace("&#x20;", "").strip()


def build():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    headings = parse_md(lines)
    doc = Document()
    configure_doc(doc)
    add_cover(doc)
    add_toc(doc, headings)

    i = 0
    diagram_no = 0
    skipped_title = False
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("# "):
            i += 1
            continue
        if line == f"## {MODULE_TITLE}" and not skipped_title:
            skipped_title = True
            i += 1
            continue
        if line.startswith("```mermaid"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            code = "\n".join(code_lines)
            diagram_no += 1
            title = f"{MODULE_NAME}架构与流程图 {diagram_no}"
            out = ASSET_DIR / f"diagram-{diagram_no:02d}.png"
            render_framework_diagram(diagram_no, out, title, code)
            add_picture(doc, out, 6.5, title)
            i += 1
            continue
        hm = re.match(r"^(##|###|####)\s+(.+)$", line)
        if hm:
            level = len(hm.group(1)) - 1
            doc.add_paragraph(hm.group(2), style=f"Heading {min(level, 3)}")
            i += 1
            continue
        if line.startswith("| "):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                parts = [p.strip() for p in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", p.replace(" ", "")) for p in parts):
                    rows.append(parts)
                i += 1
            add_table(doc, rows)
            continue
        im = re.match(r"!\[([^]]*)\]\(([^)]+)\)", line)
        if im:
            path = (SOURCE.parent / im.group(2)).resolve()
            if path.exists():
                add_picture(doc, path, 6.55, im.group(1) or path.stem)
            i += 1
            continue
        if re.match(r"^\*图\s*\d+", line):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(9)
            p.paragraph_format.keep_with_next = False
            set_run_font(p.add_run(line.strip("*")), 9, color=MUTED, italic=True)
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.12)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            set_run_font(p.add_run("说明｜" + line[2:]), 9.5, color=BLUE)
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            text = clean_inline(re.sub(r"^\d+\.\s+", "", line))
            p = doc.add_paragraph(style="List Number")
            set_run_font(p.add_run(text), 10.5)
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            set_run_font(p.add_run(clean_inline(line[2:])), 10.5)
            i += 1
            continue
        if line and line != "---":
            p = doc.add_paragraph()
            # Convert simple Markdown emphasis/code markers to clean Word text.
            clean = clean_inline(line)
            set_run_font(p.add_run(clean), 10.5)
        i += 1

    props = doc.core_properties
    props.title = f"体育博物馆智慧化综合管理平台—{MODULE_TITLE}"
    props.subject = f"{MODULE_NAME}功能架构、业务流程、操作说明与验收建议"
    props.author = "项目组"
    props.keywords = f"体育博物馆, 智慧化平台, {MODULE_NAME}, 功能说明"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
